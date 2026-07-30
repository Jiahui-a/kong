"""单元测试。"""

import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from cascon_sync.database import load_database, parse_designators
from cascon_sync.matcher import (
    discover_project_dirs,
    match_all_sources,
    match_folder_in_project,
    match_project_folders,
    parse_chip_folder_name,
)
from cascon_sync.sync import sync_projects


def _create_database_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["序号", "料号", "型号", "ProjectA", "ProjectB"])
    sheet.append(["1", "C12345-001", "STM32F103C8T6", "U1", "U5"])
    sheet.append(["2", "C12345-002", "GD32F303CCT6", "U2", ""])
    workbook.save(path)


class CasconSyncTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = Path(tempfile.mkdtemp())

    def tearDown(self) -> None:
        shutil.rmtree(self.temp_dir)

    def _load_db(self):
        xlsx = self.temp_dir / "database.xlsx"
        _create_database_xlsx(xlsx)
        return load_database(xlsx)

    def test_load_database_columns(self) -> None:
        database = self._load_db()
        self.assertEqual(len(database.records), 2)
        self.assertEqual(database.project_names, ["ProjectA", "ProjectB"])
        self.assertEqual(database.records[0].part_number, "C12345-001")
        self.assertEqual(database.records[0].model, "STM32F103C8T6")
        self.assertEqual(database.records[0].designators_in("ProjectA"), ["U1"])
        self.assertEqual(database.records[0].designators_in("ProjectB"), ["U5"])
        self.assertEqual(database.records[0].target_name("ProjectA", "U1"), "[STM32F103C8T6 C12345-001 U1]")

    def test_parse_bracket_name(self) -> None:
        self.assertEqual(parse_chip_folder_name("[U1]"), "U1")
        self.assertEqual(parse_chip_folder_name("[STM32F103C8T6 U1]"), "STM32F103C8T6 U1")
        self.assertEqual(parse_chip_folder_name("[Car_Interface]_02"), "Car_Interface")
        self.assertEqual(parse_chip_folder_name("[U1]_revA"), "U1")

    def test_match_designator_only(self) -> None:
        database = self._load_db()
        result = match_folder_in_project("[U1]", "ProjectA", database.records)
        self.assertIsNotNone(result)
        record, match_type, _, _ = result  # type: ignore[misc]
        self.assertEqual(record.part_number, "C12345-001")
        self.assertEqual(match_type, "designator_exact")

    def test_match_model_only(self) -> None:
        database = self._load_db()
        result = match_folder_in_project("[STM32F103C8T6]", "ProjectA", database.records)
        self.assertIsNotNone(result)
        record, match_type, _, _ = result  # type: ignore[misc]
        self.assertEqual(record.part_number, "C12345-001")
        self.assertEqual(match_type, "model_exact")

    def test_match_mixed_name(self) -> None:
        database = self._load_db()
        for folder_name in ("[STM32F103C8T6 U1]", "[U1 STM32F103C8T6]", "[U1-STM32F103C8T6]"):
            result = match_folder_in_project(folder_name, "ProjectA", database.records)
            self.assertIsNotNone(result, folder_name)
            record, _, _, _ = result  # type: ignore[misc]
            self.assertEqual(record.part_number, "C12345-001")

    def test_match_respects_project_designator(self) -> None:
        database = self._load_db()
        result_a = match_folder_in_project("[U5]", "ProjectA", database.records)
        result_b = match_folder_in_project("[U5]", "ProjectB", database.records)
        self.assertIsNone(result_a)
        self.assertIsNotNone(result_b)
        record, _, _, _ = result_b  # type: ignore[misc]
        self.assertEqual(record.part_number, "C12345-001")

    def test_discover_projects_from_workspace(self) -> None:
        database = self._load_db()
        workspace = self.temp_dir / "cascon"
        (workspace / "ProjectA").mkdir(parents=True)
        (workspace / "ProjectB").mkdir(parents=True)
        (workspace / "Other").mkdir(parents=True)

        projects = discover_project_dirs([workspace], database)
        names = {name for _, name in projects}
        self.assertEqual(names, {"ProjectA", "ProjectB"})

    def test_sync_with_internal_rename(self) -> None:
        database = self._load_db()
        workspace = self.temp_dir / "cascon"
        chip_folder = workspace / "ProjectA" / "[U1]"
        chip_folder.mkdir(parents=True)
        (chip_folder / "[U1]").mkdir()
        (chip_folder / "[U1].txt").write_text("data", encoding="utf-8")
        (chip_folder / "readme.txt").write_text("ok", encoding="utf-8")

        output = self.temp_dir / "database"
        report = sync_projects([workspace], database, output)

        self.assertEqual(report.copied_count, 1)
        dest = output / "[STM32F103C8T6 C12345-001 U1]"
        self.assertTrue(dest.is_dir())
        self.assertTrue((dest / "[STM32F103C8T6 C12345-001 U1]").is_dir())
        self.assertTrue((dest / "[STM32F103C8T6 C12345-001 U1].txt").is_file())
        self.assertTrue((dest / "readme.txt").is_file())

    def test_unmatched_folder_reported(self) -> None:
        database = self._load_db()
        workspace = self.temp_dir / "cascon"
        (workspace / "ProjectA" / "[UNKNOWN]").mkdir(parents=True)

        matched, unmatched, _ = match_all_sources([workspace], database)
        self.assertEqual(len(matched), 0)
        self.assertEqual(len(unmatched), 1)

    def test_parse_multiline_designators(self) -> None:
        self.assertEqual(parse_designators("U1\nU2\nU3"), ["U1", "U2", "U3"])
        self.assertEqual(parse_designators("U1,U2;U3"), ["U1", "U2", "U3"])
        self.assertEqual(parse_designators("U1、U2、U3"), ["U1", "U2", "U3"])
        self.assertEqual(parse_designators("U1\nU1"), ["U1"])

    def test_match_multiline_designators_in_cell(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-003", "W25Q128JVSIQ", "U3\nU4\nU5"])
        xlsx = self.temp_dir / "multiline.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(database.records[0].designators_in("ProjectA"), ["U3", "U4", "U5"])

        for folder in ("[U3]", "[U4]", "[U5]", "[W25Q128JVSIQ U4]"):
            result = match_folder_in_project(folder, "ProjectA", database.records)
            self.assertIsNotNone(result, folder)
            record, _, _, _ = result  # type: ignore[misc]
            self.assertEqual(record.part_number, "C12345-003")

    def test_match_project_folders(self) -> None:
        database = self._load_db()
        project = self.temp_dir / "ProjectA"
        (project / "[U1]").mkdir(parents=True)
        (project / "[U2]").mkdir(parents=True)

        matched, unmatched = match_project_folders(project, "ProjectA", database.records)
        self.assertEqual(len(matched), 2)
        self.assertEqual(len(unmatched), 0)

    def test_empty_model_allowed(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-010", "", "U10"])
        xlsx = self.temp_dir / "empty_model.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        record = database.records[0]
        self.assertEqual(record.model, "")
        self.assertEqual(record.part_number, "C12345-010")
        self.assertEqual(record.target_name("ProjectA", "U10"), "[C12345-010 U10]")

        result = match_folder_in_project("[U10]", "ProjectA", database.records)
        self.assertIsNotNone(result)
        record, match_type, _, designator = result  # type: ignore[misc]
        self.assertEqual(record.part_number, "C12345-010")
        self.assertEqual(match_type, "designator_exact")
        self.assertEqual(designator, "U10")

    def test_target_name_omits_empty_model_and_designator(self) -> None:
        from cascon_sync.database import ChipRecord

        record = ChipRecord(
            part_number="C12345-011",
            model="W25Q128",
            designators_by_project={},
            row_index=2,
        )
        self.assertEqual(record.target_name("ProjectA"), "[W25Q128 C12345-011]")

        record_no_model = ChipRecord(
            part_number="C12345-010",
            model="",
            designators_by_project={"ProjectA": ["U10"]},
            row_index=3,
        )
        self.assertEqual(record_no_model.target_name("ProjectA", "U10"), "[C12345-010 U10]")

    def test_skip_row_when_all_project_columns_empty(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "", "STM32F103C8T6", ""])
        sheet.append(["2", "C12345-020", "GD32F303", "U2"])
        xlsx = self.temp_dir / "empty_part_number.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(len(database.records), 1)
        self.assertEqual(database.records[0].part_number, "C12345-020")

    def test_row_with_empty_part_number_kept_when_project_has_value(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "", "STM32F103C8T6", "U1"])
        sheet.append(["2", "C12345-020", "GD32F303", "U2"])
        xlsx = self.temp_dir / "empty_part_with_project.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(len(database.records), 2)

        model_only = next(record for record in database.records if record.model == "STM32F103C8T6")
        self.assertEqual(model_only.part_number, "")
        self.assertEqual(model_only.designators_in("ProjectA"), ["U1"])

        result = match_folder_in_project("[U1]", "ProjectA", database.records)
        self.assertIsNotNone(result)
        record, match_type, _, _ = result  # type: ignore[misc]
        self.assertEqual(record.model, "STM32F103C8T6")
        self.assertEqual(match_type, "designator_exact")

    def test_empty_part_number_not_inherited_from_row_above(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-100", "STM32F103C8T6", "U1"])
        sheet.append(["2", "", "", "U2"])
        xlsx = self.temp_dir / "no_inherit_part_number.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(len(database.records), 2)

        with_model = next(r for r in database.records if r.model == "STM32F103C8T6")
        self.assertEqual(with_model.part_number, "C12345-100")
        self.assertEqual(with_model.designators_in("ProjectA"), ["U1"])

        without_part_number = next(r for r in database.records if r.designators_in("ProjectA") == ["U2"])
        self.assertEqual(without_part_number.part_number, "")
        self.assertEqual(without_part_number.model, "")
        self.assertEqual(without_part_number.target_name("ProjectA", "U2"), "[U2]")

        result = match_folder_in_project("[U2]", "ProjectA", database.records)
        self.assertIsNotNone(result)
        record, _, _, designator = result  # type: ignore[misc]
        self.assertEqual(record.part_number, "")
        self.assertEqual(record.model, "")
        self.assertEqual(designator, "U2")

    def test_model_not_inherited_unless_cell_has_value(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-400", "STM32F103", "U1"])
        sheet.append(["2", "C12345-401", "", "U2"])
        xlsx = self.temp_dir / "no_model_inherit.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        empty_model_rows = [r for r in database.records if r.designators_in("ProjectA") == ["U2"]]
        self.assertEqual(len(empty_model_rows), 1)
        self.assertEqual(empty_model_rows[0].part_number, "C12345-401")
        self.assertEqual(empty_model_rows[0].model, "")
        self.assertEqual(empty_model_rows[0].target_name("ProjectA", "U2"), "[C12345-401 U2]")

    def test_skip_empty_project_designator_but_keep_other_projects(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA", "ProjectB"])
        sheet.append(["1", "C12345-600", "STM32F103", "U1", ""])
        sheet.append(["2", "C12345-601", "GD32F303", "", "U9"])
        xlsx = self.temp_dir / "skip_empty_project.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(len(database.records), 2)

        record_a = next(r for r in database.records if r.part_number == "C12345-600")
        self.assertEqual(record_a.designators_in("ProjectA"), ["U1"])
        self.assertEqual(record_a.designators_in("ProjectB"), [])

        record_b = next(r for r in database.records if r.part_number == "C12345-601")
        self.assertEqual(record_b.designators_in("ProjectA"), [])
        self.assertEqual(record_b.designators_in("ProjectB"), ["U9"])

    def test_match_bracket_with_suffix(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-500", "Car_Interface", "U20"])
        xlsx = self.temp_dir / "suffix_folder.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        for folder in ("[Car_Interface]_02", "[U20]_revA", "[C12345-500]_01"):
            result = match_folder_in_project(folder, "ProjectA", database.records)
            self.assertIsNotNone(result, folder)
            record, _, parsed, _ = result  # type: ignore[misc]
            self.assertEqual(record.part_number, "C12345-500")
            self.assertNotIn("]", parsed)

    def test_match_any_single_field(self) -> None:
        database = self._load_db()
        # 料号单独命中
        result = match_folder_in_project("[C12345-001]", "ProjectA", database.records)
        self.assertIsNotNone(result)
        record, match_type, _, _ = result  # type: ignore[misc]
        self.assertEqual(record.part_number, "C12345-001")
        self.assertEqual(match_type, "part_number_exact")

        # 位号单独命中（带后缀）
        result = match_folder_in_project("[U1]_02", "ProjectA", database.records)
        self.assertIsNotNone(result)
        record, match_type, _, _ = result  # type: ignore[misc]
        self.assertEqual(match_type, "designator_exact")

    def test_merged_part_number_cells(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-200", "STM32F103", "U1、U2"])
        sheet.append(["2", None, None, "U3"])
        sheet.merge_cells("B2:B3")
        sheet.merge_cells("C2:C3")
        xlsx = self.temp_dir / "merged_part_number.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(len(database.records), 1)
        record = database.records[0]
        self.assertEqual(record.part_number, "C12345-200")
        self.assertEqual(record.model, "STM32F103")
        self.assertEqual(record.designators_in("ProjectA"), ["U1", "U2", "U3"])

        for folder in ("[U1]", "[U2]", "[U3]", "[STM32F103 U3]"):
            result = match_folder_in_project(folder, "ProjectA", database.records)
            self.assertIsNotNone(result, folder)

    def test_merged_part_number_only_model_stays_empty(self) -> None:
        """常见写法：仅合并料号，后续行只填位号；型号未合并则保持为空。"""
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-201", "STM32F103", "U1"])
        sheet.append(["2", None, None, "U2"])
        sheet.merge_cells("B2:B3")
        xlsx = self.temp_dir / "merged_part_only.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(len(database.records), 2)

        with_model = next(r for r in database.records if r.model == "STM32F103")
        self.assertEqual(with_model.part_number, "C12345-201")
        self.assertEqual(with_model.designators_in("ProjectA"), ["U1"])

        empty_model = next(r for r in database.records if r.designators_in("ProjectA") == ["U2"])
        self.assertEqual(empty_model.part_number, "C12345-201")
        self.assertEqual(empty_model.model, "")
        self.assertEqual(empty_model.target_name("ProjectA", "U2"), "[C12345-201 U2]")

    def test_merged_header_project_column(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA", None])
        sheet.append(["1", "C12345-300", "GD32F303", "U7", None])
        sheet.merge_cells("D1:E1")
        xlsx = self.temp_dir / "merged_header.xlsx"
        workbook.save(xlsx)

        database = load_database(xlsx)
        self.assertEqual(database.project_names, ["ProjectA"])
        self.assertEqual(database.records[0].designators_in("ProjectA"), ["U7"])

    def test_sync_with_empty_model(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "料号", "型号", "ProjectA"])
        sheet.append(["1", "C12345-010", "", "U10"])
        xlsx = self.temp_dir / "empty_model.xlsx"
        workbook.save(xlsx)
        database = load_database(xlsx)

        workspace = self.temp_dir / "cascon"
        chip_folder = workspace / "ProjectA" / "[U10]"
        chip_folder.mkdir(parents=True)
        (chip_folder / "[U10].txt").write_text("data", encoding="utf-8")

        output = self.temp_dir / "database"
        report = sync_projects([workspace], database, output)

        self.assertEqual(report.copied_count, 1)
        dest = output / "[C12345-010 U10]"
        self.assertTrue(dest.is_dir())
        self.assertTrue((dest / "[C12345-010 U10].txt").is_file())


if __name__ == "__main__":
    unittest.main()
